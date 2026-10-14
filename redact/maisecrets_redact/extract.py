"""Extractors (TASK-34.2): a document becomes flattened text and a map from text offsets back to its source. One
function per format. A table also keeps where each of its cells lies in the text, so its redacted cells can be written
back as CSV (TASK-54); a PDF is never rebuilt (docs/REDACT.md).

- PDF: digital text with pypdf (BSD-3-Clause), page by page. No OCR: a PDF without text raises NoText.
- CSV: the csv module, cell by cell, the delimiter sniffed; a column can be dropped by its header.
- Excel: openpyxl (MIT), every sheet, cell values (not formulas), a column dropped by its header.
- .txt and .md pass through.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from .pipeline import Document, Position, Table


class Unsupported(ValueError):
    """A file type the tool does not extract."""


class NoText(ValueError):
    """A document with no extractable text (a scanned PDF: OCR is out of scope)."""


class _Builder:
    """Text and positions built together, so every offset maps back to where it came from."""

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.positions: list[Position] = []
        self.tables: list[Table] = []
        self.size = 0

    def add(self, text: str, where: str | None = None) -> None:
        if where is not None and text:
            self.positions.append(Position(self.size, self.size + len(text), where))
        self.parts.append(text)
        self.size += len(text)

    def document(self, path: Path) -> Document:
        return Document(str(path), "".join(self.parts), tuple(self.positions), tuple(self.tables))


def _drop(headers: list[str], drop: set[str]) -> set[int]:
    wanted = {d.strip().lower() for d in drop}
    return {i for i, h in enumerate(headers) if str(h or "").strip().lower() in wanted}


def _table(b: _Builder, rows: list[list[str]], drop: set[str], prefix: str = "", name: str | None = None,
           delimiter: str = ",", name_at: tuple[int, int] | None = None) -> None:
    """Rows as tab-separated lines; the first row names the columns. The table, without its dropped columns, is
    recorded too."""
    headers = [str(h) if h is not None else "" for h in rows[0]] if rows else []
    skip = _drop(headers, drop)
    table = []
    for r, row in enumerate(rows, 1):
        cells = [(c, v) for c, v in enumerate(row) if c not in skip]
        spans = []
        for n, (c, value) in enumerate(cells):
            header = headers[c] if c < len(headers) and headers[c] else _column_letter(c)
            start = b.size
            b.add("" if value is None else str(value), f"{prefix}row {r}, column {header}")
            spans.append((start, b.size))
            if n < len(cells) - 1:
                b.add("\t")
        b.add("\n")
        table.append(tuple(spans))
    b.tables.append(Table(name, delimiter, tuple(table), name_at))


def _column_letter(i: int) -> str:
    out = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        out = chr(65 + rem) + out
    return out


def _pdf(path: Path) -> Document:
    from pypdf import PdfReader
    b = _Builder()
    for n, page in enumerate(PdfReader(str(path)).pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            b.add(text, f"page {n}")
            b.add("\n\n")
    doc = b.document(path)
    if not doc.text.strip():
        raise NoText(f"{path.name}: no text found (a scanned PDF? OCR is out of scope)")
    return doc


def _csv(path: Path, drop: set[str]) -> Document:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    b = _Builder()
    # newline="": a quoted cell may hold a line break
    _table(b, list(csv.reader(io.StringIO(raw, newline=""), dialect)), drop, delimiter=dialect.delimiter)
    return b.document(path)


def _excel(path: Path, drop: set[str]) -> Document:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    b = _Builder()
    try:
        for ws in wb.worksheets:
            at = b.size + 1
            b.add(f"[{ws.title}]\n")
            _table(b, [list(r) for r in ws.iter_rows(values_only=True)], drop, prefix=f"sheet {ws.title}, ",
                   name=ws.title, name_at=(at, at + len(ws.title)))
            b.add("\n")
    finally:
        wb.close()
    return b.document(path)


def _plain(path: Path) -> Document:
    b = _Builder()
    b.add(path.read_text(encoding="utf-8", errors="replace"), "file")
    return b.document(path)


SUPPORTED = (".pdf", ".csv", ".tsv", ".xlsx", ".xlsm", ".txt", ".md")


def extract(path: Path | str, drop_columns: set[str] | frozenset[str] = frozenset()) -> Document:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix in (".csv", ".tsv"):
        return _csv(path, set(drop_columns))
    if suffix in (".xlsx", ".xlsm"):
        return _excel(path, set(drop_columns))
    if suffix in (".txt", ".md"):
        return _plain(path)
    raise Unsupported(f"{path.name}: {suffix.lstrip('.') or 'no extension'} is not a supported type "
                      f"({', '.join(SUPPORTED)})")
