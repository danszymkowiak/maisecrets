"""A table comes out as a table (TASK-54, user 2026-10-14): a CSV or TSV as <file>.redacted.csv with its delimiter and
its header row, a workbook as one <file>.<sheet>.redacted.csv per sheet. Tab-separated text without quoting shifted a
row's columns when a cell held a tab or a line break; since TASK-49 no span crosses a cell, so each cell is redacted
on its own and the rows are written back as CSV.

Run: uv run --project redact python -m unittest discover -s redact/tests
"""
from __future__ import annotations

import csv
import io
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _home  # noqa: E402,F401  first: a temp maisecrets home

from maisecrets_redact.adapters import compose as app  # noqa: E402

from maisecrets_redact.domain import text as text_ops  # noqa: E402
from maisecrets_redact.domain.document import Span  # noqa: E402

MAIL = "max@mailbox-test.de"


def rows(path: Path, delimiter: str = ",") -> list[list[str]]:
    return list(csv.reader(io.StringIO(path.read_text(encoding="utf-8"), newline=""), delimiter=delimiter))


def run(src: Path, **kw) -> app.Result:
    return app.run_file(src, src.with_name(app.output_name(src)), model_detectors={}, regions=("generic",),
                        interactive=False, **kw)


class ToCsvTests(unittest.TestCase):
    def test_cells_with_a_delimiter_quote_tab_or_line_break_are_quoted(self):
        text = text_ops.to_csv([["a", "b,c", 'say "hi"', "x\ty", "two\nlines", ""]], ",")
        self.assertEqual(text, 'a,"b,c","say ""hi""","x\ty","two\nlines",\r\n')
        self.assertEqual(list(csv.reader(io.StringIO(text, newline=""))), [["a", "b,c", 'say "hi"', "x\ty",
                                                                            "two\nlines", ""]])

    def test_the_delimiter_of_the_input(self):
        self.assertEqual(text_ops.to_csv([["a;b", "c"]], ";"), '"a;b";c\r\n')


class SafeNameTests(unittest.TestCase):
    def test_a_sheet_name_unsafe_in_a_file_name_is_made_safe_and_names_stay_apart(self):
        self.assertEqual(text_ops.safe_names(["Lohn 2025", "a/b", "a:b", "..", "", "Konten"]),
                         ["Lohn 2025", "a_b", "a_b-2", "sheet4", "sheet5", "Konten"])


class CsvOutputTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_a_csv_keeps_its_delimiter_header_and_quoted_cells(self):
        src = self.dir / "people.csv"
        src.write_text(f'name;email;note\nMax;{MAIL};"a; b\nwith a break"\n', encoding="utf-8")
        self.assertEqual(run(src).status, "written")
        out = self.dir / "people.csv.redacted.csv"
        self.assertEqual(rows(out, ";"), [["name", "email", "note"], ["Max", "⟦EMAIL_1⟧", "a; b\nwith a break"]])
        self.assertFalse((self.dir / "people.csv.redacted.txt").exists())

    def test_a_tsv_stays_tab_separated(self):
        src = self.dir / "people.tsv"
        src.write_text(f"name\temail\nMax\t{MAIL}\n", encoding="utf-8")
        run(src)
        self.assertEqual(rows(self.dir / "people.tsv.redacted.csv", "\t"), [["name", "email"], ["Max", "⟦EMAIL_1⟧"]])

    def test_a_model_span_in_a_cell_and_a_dropped_column(self):
        src = self.dir / "people.csv"
        src.write_text(f"name,email,iban\nMax Muster,{MAIL},x\n", encoding="utf-8")

        def model(name):
            def find(doc):
                at = doc.text.index("Max Muster")
                return [Span(at, at + 10, "NAME", name)]
            return find
        app.run_file(src, src.with_name(app.output_name(src)), model_detectors={"m1": model("m1"), "m2": model("m2")},
                     regions=("generic",), interactive=False, drop_columns=frozenset({"iban"}))
        self.assertEqual(rows(self.dir / "people.csv.redacted.csv"), [["name", "email"], ["⟦NAME_1⟧", "⟦EMAIL_1⟧"]])

    def test_a_workbook_gives_one_csv_per_sheet_and_one_value_keeps_one_placeholder(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Lohn|2025"
        wb.active.append(["email", "note"])
        wb.active.append([MAIL, "a,b"])
        other = wb.create_sheet("Konten")
        other.append(["contact"])
        other.append([MAIL])
        src = self.dir / "lohn.xlsx"
        wb.save(src)
        self.assertEqual(app.output_name(src), "lohn.xlsx.redacted.csv")
        result = run(src)
        self.assertEqual(result.status, "written")
        self.assertEqual(sorted(p.name for p in self.dir.glob("*.redacted.csv")),
                         ["lohn.xlsx.Konten.redacted.csv", "lohn.xlsx.Lohn_2025.redacted.csv"])
        self.assertEqual(rows(self.dir / "lohn.xlsx.Lohn_2025.redacted.csv"), [["email", "note"], ["⟦EMAIL_1⟧", "a,b"]])
        self.assertEqual(rows(self.dir / "lohn.xlsx.Konten.redacted.csv"), [["contact"], ["⟦EMAIL_1⟧"]])

    def test_a_sheet_whose_name_is_redacted_is_named_by_its_number(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = MAIL
        wb.active.append(["x"])
        src = self.dir / "w.xlsx"
        wb.save(src)
        run(src)
        self.assertEqual([p.name for p in self.dir.glob("*.redacted.csv")], ["w.xlsx.sheet1.redacted.csv"])

    def test_an_existing_sheet_output_is_never_overwritten(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "A"
        wb.active.append([MAIL])
        wb.create_sheet("B").append(["x"])
        src = self.dir / "w.xlsx"
        wb.save(src)
        (self.dir / "w.xlsx.B.redacted.csv").write_text("mine\n", encoding="utf-8")
        self.assertEqual(run(src).status, "exists")
        self.assertEqual((self.dir / "w.xlsx.B.redacted.csv").read_text(encoding="utf-8"), "mine\n")
        self.assertFalse((self.dir / "w.xlsx.A.redacted.csv").exists(), "nothing of the workbook is written")

    def test_text_documents_still_come_out_as_text(self):
        src = self.dir / "notes.md"
        self.assertEqual(app.output_name(src), "notes.md.redacted.txt")


class DirectoryTests(unittest.TestCase):
    def test_a_second_run_skips_every_kind_of_output(self):
        d = Path(tempfile.mkdtemp())
        (d / "a.csv").write_text(f"email\n{MAIL}\n", encoding="utf-8")
        (d / "b.txt").write_text(f"mail {MAIL}\n", encoding="utf-8")
        first = app.run_dir(d, d, model_detectors={}, regions=("generic",), interactive=False)
        self.assertEqual(first.written, 2)
        self.assertTrue((d / "a.csv.redacted.csv").exists())
        again = app.run_dir(d, d, model_detectors={}, regions=("generic",), interactive=False)
        self.assertEqual((again.written, again.exists), (0, 2))
        self.assertFalse(list(d.glob("*.redacted.csv.redacted.*")))


if __name__ == "__main__":
    unittest.main()
