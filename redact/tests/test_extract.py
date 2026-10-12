"""Extractors (TASK-34.2): each format becomes flattened text plus a map from text offsets back to the source, so a
review can say where a span is. Fixtures are generated here with made-up values; no binary file sits in the tree.

Run: uv run --project redact python -m unittest discover -s redact/tests
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from maisecrets_redact import extract


def minimal_pdf(pages: list[list[str]]) -> bytes:
    """A valid PDF with one Helvetica text line per string on each page."""
    objs: list[bytes] = []
    kids = []
    font_id = 3
    objs.append(b"")                                   # 1: catalog, filled below
    objs.append(b"")                                   # 2: pages, filled below
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for lines in pages:
        ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
        for line in lines:
            esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops += [f"({esc}) Tj", "T*"]
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_id = len(objs)
        objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
                    b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (content_id, font_id))
        kids.append(len(objs))
    objs[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


class ExtractTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def write(self, name: str, data) -> Path:
        p = self.dir / name
        if isinstance(data, bytes):
            p.write_bytes(data)
        else:
            p.write_text(data, encoding="utf-8")
        return p

    def where(self, doc, needle: str) -> str:
        at = doc.text.index(needle)
        return next(p.where for p in doc.positions if p.start <= at < p.end)

    def test_a_pdf_gives_its_text_and_the_page_of_each_line(self):
        p = self.write("statement.pdf", minimal_pdf([["Kontoauszug 2026", "IBAN CH93 0076 2011 6238 5295 7"],
                                                    ["Seite zwei", "Mail max@firma.example"]]))
        doc = extract.extract(p)
        self.assertIn("CH93 0076 2011 6238 5295 7", doc.text)
        self.assertEqual(self.where(doc, "CH93"), "page 1")
        self.assertEqual(self.where(doc, "max@firma"), "page 2")

    def test_a_pdf_without_text_says_so(self):
        p = self.write("scan.pdf", minimal_pdf([[]]))
        with self.assertRaisesRegex(extract.NoText, "no text"):
            extract.extract(p)

    def test_a_csv_gives_each_cell_with_its_row_and_column(self):
        p = self.write("people.csv", "name;email;note\nMax Muster;max@firma.example;ok\nEva;eva@firma.example;-\n")
        doc = extract.extract(p)
        self.assertIn("max@firma.example", doc.text)
        self.assertEqual(self.where(doc, "eva@firma"), "row 3, column email")

    def test_a_csv_column_can_be_dropped_by_its_header(self):
        p = self.write("people.csv", "name,email,note\nMax Muster,max@firma.example,ok\n")
        doc = extract.extract(p, drop_columns={"EMAIL"})
        self.assertNotIn("max@firma.example", doc.text)
        self.assertIn("Max Muster", doc.text)

    def test_an_excel_workbook_gives_values_of_every_sheet(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Lohn"
        ws.append(["Name", "AHV", "Betrag"])
        ws.append(["Max Muster", "756.9217.0769.85", 5200])
        ws["D2"] = "=C2*12"
        other = wb.create_sheet("Konten")
        other.append(["IBAN"])
        other.append(["CH93 0076 2011 6238 5295 7"])
        p = self.dir / "lohn.xlsx"
        wb.save(p)
        doc = extract.extract(p)
        self.assertIn("756.9217.0769.85", doc.text)
        self.assertEqual(self.where(doc, "756.9217"), "sheet Lohn, row 2, column AHV")
        self.assertEqual(self.where(doc, "CH93"), "sheet Konten, row 2, column IBAN")
        self.assertNotIn("=C2*12", doc.text, "values, not formulas")

    def test_an_excel_column_can_be_dropped_by_its_header(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["Name", "AHV"])
        wb.active.append(["Max Muster", "756.9217.0769.85"])
        p = self.dir / "lohn.xlsx"
        wb.save(p)
        self.assertNotIn("756.9217", extract.extract(p, drop_columns={"ahv"}).text)

    def test_plain_text_passes_through(self):
        p = self.write("notes.md", "# Notes\npassword: see vault\n")
        doc = extract.extract(p)
        self.assertEqual(doc.text, "# Notes\npassword: see vault\n")
        self.assertEqual(self.where(doc, "password"), "file")

    def test_an_unsupported_type_is_named(self):
        p = self.write("photo.jpg", b"\xff\xd8\xff")
        with self.assertRaisesRegex(extract.Unsupported, "jpg"):
            extract.extract(p)


if __name__ == "__main__":
    unittest.main()
