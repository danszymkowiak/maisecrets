"""The deterministic pass (TASK-34.3): the plugin's rules find secrets and structured personal data in the extracted
text of every format, each value becomes a placeholder, the same value the same one across a run, and the run leaves
the plugin's home untouched.

Run: uv run --project redact python -m unittest discover -s redact/tests
"""
from __future__ import annotations

import os
import random
import string
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _home  # noqa: E402  first: a temp maisecrets home

from maisecrets_redact.adapters import extract  # noqa: E402

from maisecrets_redact.domain import document  # noqa: E402

from maisecrets_redact.domain import text as text_ops  # noqa: E402

from maisecrets_redact.domain import placeholders  # noqa: E402

from maisecrets_redact.adapters import plugin_rules  # noqa: E402
from test_extract import minimal_pdf  # noqa: E402

_R = random.Random(343)
TOKEN = "ghp_" + "".join(_R.choice(string.ascii_letters + string.digits) for _ in range(36))
IBAN = "CH93 0076 2011 6238 5295 7"            # the registry's Swiss example
CARD = "4111 1111 1111 1111"                   # the common Luhn test number
MAIL = "max.muster@firma-beispiel.ch"
AHV = "756.9217.0769.85"                       # the ZAS example


def redact_text(doc, regions=("generic",)):
    det = plugin_rules.detector(regions=regions)
    names = placeholders.Placeholders()
    return text_ops.render(doc, document.detect(doc, [det]), names.placeholder), names


class RulesTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_each_format_comes_out_with_placeholders(self):
        line = f"token {TOKEN} iban {IBAN} card {CARD} mail {MAIL}"
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["what", "value"])
        wb.active.append(["line", line])
        wb.save(self.dir / "a.xlsx")
        (self.dir / "a.txt").write_text(line + "\n", encoding="utf-8")
        (self.dir / "a.csv").write_text(f"what,value\nline,{line}\n", encoding="utf-8")
        (self.dir / "a.pdf").write_bytes(minimal_pdf([[line]]))
        for name in ("a.txt", "a.csv", "a.xlsx", "a.pdf"):
            with self.subTest(name=name):
                out, _ = redact_text(extract.extract(self.dir / name))
                for value in (TOKEN, IBAN, CARD, MAIL):
                    self.assertNotIn(value, out)
                for kind in ("SECRET", "IBAN", "CARD", "EMAIL"):
                    self.assertIn(f"⟦{kind}_1⟧", out)

    def test_the_same_value_gets_the_same_placeholder_and_another_value_another(self):
        other = "eva@firma-beispiel.ch"
        doc = document.Document("x.txt", f"{MAIL}, {other} und nochmals {MAIL}")
        out, names = redact_text(doc)
        self.assertEqual(out, "⟦EMAIL_1⟧, ⟦EMAIL_2⟧ und nochmals ⟦EMAIL_1⟧")
        self.assertEqual(names.counts(), {"EMAIL": 2})

    def test_regions_choose_the_country_rules(self):
        doc = document.Document("x.txt", f"AHV-Nr. {AHV}")
        self.assertIn(AHV, redact_text(doc)[0], "generic rules only: no Swiss number")
        self.assertNotIn(AHV, redact_text(doc, regions=("generic", "ch"))[0])

    def test_a_run_leaves_the_plugin_home_untouched(self):
        before = sorted(os.listdir(_home.HOME))
        redact_text(document.Document("x.txt", f"token {TOKEN}"))
        self.assertEqual(sorted(os.listdir(_home.HOME)), before)


class ModelsAfterRulesTests(unittest.TestCase):
    """The models run after the rules (user, 2026-10-12): they never read a value the rules found."""

    def test_a_model_reads_the_text_with_rule_hits_replaced_and_its_spans_land_on_the_original(self):
        from maisecrets_redact.adapters import compose as app
        seen = []

        def model(doc):
            seen.append(doc.text)
            at = doc.text.index("Erika Muster")
            return [document.Span(at, at + 12, "NAME", "m")]
        d = Path(tempfile.mkdtemp())
        src = d / "a.txt"
        src.write_text(f"deploy key {TOKEN} for Erika Muster, IBAN {IBAN}\n", encoding="utf-8")
        result = app.run_file(src, d / "a.txt.redacted.txt", {"m": model}, ("ch",), interactive=False)
        self.assertNotIn(TOKEN, seen[0])
        self.assertNotIn(IBAN, seen[0])
        self.assertIn("[SECRET]", seen[0])
        out = (d / "a.txt.redacted.txt").read_text(encoding="utf-8")
        self.assertEqual((result.status, "Erika" in out), ("written", False))


class CommandTests(unittest.TestCase):
    def run_cli(self, *args):
        env = {**os.environ, "MAISECRETS_HOME": _home.HOME}
        return subprocess.run([sys.executable, "-m", "maisecrets_redact", *args], capture_output=True, text=True,
                              timeout=120, env=env)

    def test_a_file_is_redacted_next_to_itself_and_never_overwritten(self):
        d = Path(tempfile.mkdtemp())
        src = d / "notes.txt"
        src.write_text(f"deploy key {TOKEN}\nIBAN {IBAN}\n", encoding="utf-8")
        r = self.run_cli(str(src), "--rules-only")   # the models are covered with fakes; this test needs no download
        self.assertEqual(r.returncode, 0, r.stderr)
        out = (d / "notes.txt.redacted.txt").read_text(encoding="utf-8")
        self.assertNotIn(TOKEN, out)
        self.assertIn("⟦SECRET_1⟧", out)
        self.assertEqual(src.read_text(encoding="utf-8"), f"deploy key {TOKEN}\nIBAN {IBAN}\n",
                         "the input is unchanged")
        self.assertNotIn(TOKEN, r.stdout + r.stderr, "the summary names counts, never a value")
        r = self.run_cli(str(src), "--rules-only")
        self.assertNotEqual(r.returncode, 0, "an existing output is not overwritten")
        self.assertIn("exists", r.stderr)

    @unittest.skipUnless(__import__("maisecrets_redact.adapters.model_backends").adapters.model_backends.unavailable_reason(),
                         "the model extra is installed")
    def test_without_the_model_extra_the_rules_run_and_it_says_so_once(self):
        d = Path(tempfile.mkdtemp())
        src = d / "notes.txt"
        src.write_text(f"deploy key {TOKEN}\n", encoding="utf-8")
        r = self.run_cli(str(src))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(r.stderr.count("only the rules run"), 1)
        self.assertNotIn(TOKEN, (d / "notes.txt.redacted.txt").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
